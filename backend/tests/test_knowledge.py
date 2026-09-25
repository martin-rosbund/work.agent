import pytest
from fastapi import HTTPException
from sqlalchemy import select
from app.services import save_knowledge, export_knowledge
from app.models import KnowledgeVersion, Chunk
from app.config import DATA_DIR
from app.extract import sections, extract


def test_versions_conflict_restore_and_export(db):
    note = save_knowledge(db, "A", "first")
    db.commit()
    save_knowledge(db, "B", "second", note.id, 1)
    db.commit()
    with pytest.raises(HTTPException):
        save_knowledge(db, "Wrong", "stale", note.id, 1)
    old = db.scalar(
        select(KnowledgeVersion).where(
            KnowledgeVersion.item_id == note.id, KnowledgeVersion.version == 1
        )
    )
    save_knowledge(db, old.title, old.content, note.id, 2)
    db.commit()
    export_knowledge(db, note.id)
    assert note.version == 3
    text = (DATA_DIR / "knowledge" / f"{note.id}.md").read_text()
    assert "version: 3" in text and "first" in text
    assert db.scalar(select(Chunk).where(Chunk.item_id == note.id)).text == "first"


def test_text_upload_and_limit(logged):
    response = logged.post(
        "/api/v1/documents", files={"file": ("../a.md", b"# Hello", "text/markdown")}
    )
    assert response.status_code == 200, response.text
    assert response.json()["title"] == "a.md" and "file_path" not in response.json()
    assert (
        logged.post(
            "/api/v1/documents", files={"file": ("a.exe", b"not executable")}
        ).status_code
        == 415
    )


def test_office_extraction_locators(tmp_path):
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation

    doc = Document()
    doc.add_paragraph("Ein wichtiger Beschluss")
    path = tmp_path / "x.docx"
    doc.save(path)
    assert "Beschluss" in sections(path)[0][1]
    book = Workbook()
    book.active["B3"] = "Budget 500"
    path = tmp_path / "x.xlsx"
    book.save(path)
    result = sections(path)
    assert result[0][0].endswith("Zeile 3") and "B3: Budget 500" in result[0][1]
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[0])
    slide.shapes.title.text = "Projekt"
    path = tmp_path / "x.pptx"
    deck.save(path)
    assert sections(path)[0] == ("Folie 1", "Projekt\n")


def test_encrypted_pdf_rejected(tmp_path):
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(100, 100)
    writer.encrypt("password")
    path = tmp_path / "secret.pdf"
    with path.open("wb") as f:
        writer.write(f)
    with pytest.raises(ValueError, match="Passwort"):
        sections(path)


def test_subprocess_text_extraction(tmp_path):
    path = tmp_path / "hello.md"
    path.write_text("Hallo Welt", encoding="utf-8")
    assert extract(path) == [["Text", "Hallo Welt"]]


def test_ocr_image_and_scanned_pdf(tmp_path):
    import shutil

    if not shutil.which("tesseract") or not shutil.which("pdftoppm"):
        pytest.skip("OCR binaries supplied by the production container")
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (1300, 250), "white")
    draw = ImageDraw.Draw(image)
    font = (
        ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 50)
        if __import__("pathlib")
        .Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
        .exists()
        else ImageFont.load_default(size=50)
    )
    draw.text((50, 80), "Projekt Alpha Freigabe 2026", fill="black", font=font)
    image_path = tmp_path / "scan.png"
    image.save(image_path)
    pdf_path = tmp_path / "scan.pdf"
    image.save(pdf_path, "PDF", resolution=150)
    assert "Alpha" in sections(image_path)[0][1]
    assert "Alpha" in sections(pdf_path)[0][1]


def test_citation_resolves_exact_location_and_guards_document_version(logged, db, mail):
    from app.services import replace_chunks

    _, item = mail
    item.kind = "document"
    replace_chunks(db, item, [("Seite 1", "Introduction"), ("Seite 2", "The decision")])
    db.commit()
    response = logged.get(
        f"/api/v1/items/{item.id}/citation",
        params={"locator": "Seite 2 · Abschnitt 1", "version": 1},
    )
    assert response.status_code == 200 and response.json()["text"] == "The decision"
    item.version = 2
    db.commit()
    assert (
        logged.get(
            f"/api/v1/items/{item.id}/citation", params={"version": 1}
        ).status_code
        == 409
    )


def test_knowledge_citation_can_show_historical_version(logged, db):
    note = save_knowledge(db, "Rule", "Original rule")
    db.commit()
    save_knowledge(db, "Rule", "Updated rule", note.id, 1)
    db.commit()
    result = logged.get(f"/api/v1/items/{note.id}/citation", params={"version": 1})
    assert result.status_code == 200 and result.json()["text"] == "Original rule"
