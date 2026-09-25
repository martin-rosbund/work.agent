"""Local document parsing. Runs in a subprocess with a timeout, never executes document macros."""

import json
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from app.config import MAX_FILE_BYTES, SUPPORTED_EXTENSIONS


def extract(path):
    path = Path(path)
    if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise ValueError("Dateiformat wird nicht unterstützt.")
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ValueError("Datei überschreitet 25 MB.")
    result = subprocess.run(
        [sys.executable, "-m", "app.features.knowledge.extraction", str(path)],
        capture_output=True,
        timeout=180,
        text=True,
        encoding="utf-8",
    )
    if result.returncode:
        raise ValueError(
            result.stderr.strip()[-400:] or "Dokument konnte nicht gelesen werden."
        )
    return json.loads(result.stdout)


def office_guard(path):
    with zipfile.ZipFile(path) as archive:
        if (
            sum(i.file_size for i in archive.infolist()) > 150 * 1024 * 1024
            or len(archive.infolist()) > 20000
        ):
            raise ValueError("Entpacktes Dokument ist zu groß.")


def sections(path):
    ext = path.suffix.lower()
    if ext in {".txt", ".md"}:
        return [("Text", path.read_text(encoding="utf-8-sig", errors="replace"))]
    if ext in {".docx", ".xlsx", ".pptx"}:
        office_guard(path)
    if ext == ".docx":
        from docx import Document

        doc = Document(path)
        text = "\n".join(p.text for p in doc.paragraphs)
        text += "\n" + "\n".join(
            " | ".join(cell.text for cell in row.cells)
            for table in doc.tables
            for row in table.rows
        )
        return [("Dokument", text)]
    if ext == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        result = []
        for sheet in workbook:
            for index, row in enumerate(sheet.iter_rows(), 1):
                if index > 50000:
                    raise ValueError("Tabelle überschreitet 50.000 Zeilen.")
                cells = [
                    f"{cell.coordinate}: {cell.value}"
                    for cell in row
                    if cell.value is not None
                ]
                if cells:
                    result.append((f"{sheet.title}, Zeile {index}", " | ".join(cells)))
        workbook.close()
        return result
    if ext == ".pptx":
        from pptx import Presentation

        return [
            (
                f"Folie {n}",
                "\n".join(
                    shape.text
                    if shape.has_text_frame
                    else "\n".join(
                        " | ".join(c.text for c in r.cells) for r in shape.table.rows
                    )
                    if shape.has_table
                    else ""
                    for shape in slide.shapes
                ),
            )
            for n, slide in enumerate(Presentation(path).slides, 1)
        ]
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ValueError(
                "Passwortgeschütztes PDF: bitte eine entschlüsselte Kopie hochladen."
            )
        if len(reader.pages) > 300:
            raise ValueError("PDF überschreitet 300 Seiten.")
        result = []
        for n, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            if len(text.strip()) < 30:
                with tempfile.TemporaryDirectory() as tmp:
                    base = str(Path(tmp) / "page")
                    subprocess.run(
                        [
                            "pdftoppm",
                            "-f",
                            str(n),
                            "-l",
                            str(n),
                            "-r",
                            "160",
                            "-singlefile",
                            "-png",
                            str(path),
                            base,
                        ],
                        check=True,
                        capture_output=True,
                        timeout=30,
                    )
                    text = ocr(Path(base + ".png"))
            result.append((f"Seite {n}", text))
        return result
    return [("Bild", ocr(path))]


def ocr(path):
    import pytesseract
    from PIL import Image, ImageSequence

    Image.MAX_IMAGE_PIXELS = 40_000_000
    with Image.open(path) as image:
        return "\n".join(
            pytesseract.image_to_string(frame, lang="deu+eng", timeout=40)
            for frame in ImageSequence.Iterator(image)
        )


if __name__ == "__main__":
    try:
        if sys.platform != "win32":
            import resource

            resource.setrlimit(resource.RLIMIT_AS, (1536 * 1024**2, 1536 * 1024**2))
            resource.setrlimit(resource.RLIMIT_CPU, (160, 160))
        result = sections(Path(sys.argv[1]))
        if sum(len(text) for _, text in result) > 5_000_000:
            raise ValueError("Extrahierter Text überschreitet die Verarbeitungsgrenze.")
        print(json.dumps(result, ensure_ascii=False))
    except Exception as exc:
        print(f"{type(exc).__name__}: {exc}", file=sys.stderr)
        sys.exit(1)
